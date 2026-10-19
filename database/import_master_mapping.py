
import os
import re
import sys
from pathlib import Path
from urllib.parse import urlparse, unquote

import pymysql
from dotenv import load_dotenv
from openpyxl import load_workbook


# Lokasi project dan master mapping
BASE_DIR = Path(__file__).resolve().parent.parent
MASTER_FILE = (
    BASE_DIR
    / "database"
    / "master" 
    / "Master_Mapping_Customer_Operator.xlsx"
)

# Customer yang belum dikonfirmasi: jangan aktifkan dahulu
EXCLUDED_CUSTOMERS = {
    "fobsetu",
    "dataprsmmscfd",
    "mskht",
}

VALID_OPERATORS = {"SETU", "KHT", "BLORA", "WUNUT"}


def normalize(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def connect_database():
    load_dotenv(BASE_DIR / "backend" / ".env")

    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL tidak ditemukan di backend/.env")

    parsed = urlparse(database_url)

    return pymysql.connect(
        host=parsed.hostname or "127.0.0.1",
        port=parsed.port or 3306,
        user=unquote(parsed.username or ""),
        password=unquote(parsed.password or ""),
        database=parsed.path.lstrip("/"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        autocommit=False,
    )


def read_master():
    if not MASTER_FILE.is_file():
        raise FileNotFoundError(
            f"File master mapping tidak ditemukan:\n{MASTER_FILE}"
        )

    workbook = load_workbook(MASTER_FILE, data_only=True, read_only=True)

    if "MASTER_MAPPING" not in workbook.sheetnames:
        raise ValueError("Sheet MASTER_MAPPING tidak ditemukan.")

    sheet = workbook["MASTER_MAPPING"]
    rows = sheet.iter_rows(values_only=True)

    headers = next(rows, None)
    if not headers:
        raise ValueError("Sheet MASTER_MAPPING kosong.")

    header_map = {
        normalize(value): index
        for index, value in enumerate(headers)
        if value is not None
    }

    # Cari kolom berdasarkan nama header
    customer_col = next(
        (
            header_map[key]
            for key in (
                "customer",
                "customername",
                "namacustomer",
                "customeraslimaster",
            )
            if key in header_map
        ),
        None,
    )

    operator_col = next(
        (
            header_map[key]
            for key in ("operator", "wilayah")
            if key in header_map
        ),
        None,
    )

    if customer_col is None or operator_col is None:
        workbook.close()
        raise ValueError(
            "Kolom customer/operator tidak ditemukan. "
            f"Header yang terbaca: {list(headers)}"
        )

    mappings = {}
    excluded_count = 0

    for row_number, row in enumerate(rows, start=2):
        customer = str(row[customer_col] or "").strip()
        operator = str(row[operator_col] or "").strip().upper()

        if not customer:
            continue

        key = normalize(customer)

        if key in EXCLUDED_CUSTOMERS:
            excluded_count += 1
            continue

        if operator not in VALID_OPERATORS:
            raise ValueError(
                f"Operator tidak valid pada baris {row_number}: "
                f"{customer} -> {operator}"
            )

        if key in mappings:
            previous = mappings[key]
            if previous["operator"] != operator:
                raise ValueError(
                    f"Customer duplikat dengan operator berbeda: {customer}"
                )
            raise ValueError(f"Customer duplikat di master: {customer}")

        mappings[key] = {
            "customer": customer,
            "operator": operator,
        }

    workbook.close()
    return mappings, excluded_count


def main():
    mappings, excluded_count = read_master()

    if not mappings:
        raise RuntimeError("Tidak ada mapping aktif untuk diimpor.")

    connection = connect_database()

    try:
        with connection.cursor() as cursor:
            # Periksa semua konflik terlebih dahulu sebelum menyimpan.
            conflicts = []

            for key, item in mappings.items():
                cursor.execute(
                    "SELECT customer_name, operator "
                    "FROM customers WHERE customer_key=%s",
                    (key,),
                )
                existing = cursor.fetchone()

                if existing and existing["operator"] != item["operator"]:
                    conflicts.append(
                        f'{item["customer"]}: database={existing["operator"]}, '
                        f'master={item["operator"]}'
                    )

            if conflicts:
                raise RuntimeError(
                    "Import dibatalkan karena ada mapping yang berbeda:\n"
                    + "\n".join(conflicts)
                    + "\nPeriksa mapping tersebut terlebih dahulu."
                )

            inserted = 0
            updated = 0

            for key, item in mappings.items():
                cursor.execute(
                    "SELECT customer_name FROM customers "
                    "WHERE customer_key=%s",
                    (key,),
                )
                existing = cursor.fetchone()

                if existing:
                    cursor.execute(
                        "UPDATE customers SET customer_name=%s "
                        "WHERE customer_key=%s",
                        (item["customer"], key),
                    )
                    updated += 1
                else:
                    cursor.execute(
                        "INSERT INTO customers "
                        "(customer_key, customer_name, operator) "
                        "VALUES (%s, %s, %s)",
                        (key, item["customer"], item["operator"]),
                    )
                    inserted += 1

        connection.commit()

        print("IMPORT MASTER MAPPING BERHASIL")
        print(f"Mapping aktif dalam file : {len(mappings)}")
        print(f"Customer baru disimpan   : {inserted}")
        print(f"Customer sudah ada       : {updated}")
        print(f"Customer dilewati        : {excluded_count}")
        print("Customer yang dilewati tidak diaktifkan.")

    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"IMPORT GAGAL: {error}")
        sys.exit(1)
