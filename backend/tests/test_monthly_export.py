import asyncio
import sys
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main


class MemoryMonthlyExport:
    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def execute(self, statement, parameters=None):
        if "FROM sales s" in statement:
            if parameters != ("2026-10-01", "2026-10-31"):
                raise AssertionError(f"Unexpected export date range: {parameters}")
            return SimpleNamespace(fetchall=lambda: [
                {
                    "report_date": date(2026, 10, 5),
                    "customer_name": "Customer FOB",
                    "operator": "SETU",
                    "so": 100,
                    "actual": 120,
                },
                {
                    "report_date": date(2026, 10, 6),
                    "customer_name": "Customer KHT",
                    "operator": "KHT",
                    "so": 200,
                    "actual": 180,
                },
                {
                    "report_date": date(2026, 10, 6),
                    "customer_name": "Customer OFF",
                    "operator": "OFF",
                    "so": 50,
                    "actual": 40,
                },
            ])
        if statement.startswith("SELECT so FROM monthly_so_targets"):
            return SimpleNamespace(fetchone=lambda: {"so": 1000})
        raise AssertionError(f"Unexpected SQL in export test: {statement}")


class MonthlyExportTests(unittest.TestCase):
    def test_export_creates_requested_sheets_and_month_totals(self):
        with patch.object(main, "db", return_value=MemoryMonthlyExport()):
            response = main.export_monthly_report(
                SimpleNamespace(session={"username": main.APP_USERNAME}),
                year=2026,
                month=10,
            )

        async def read_body():
            return b"".join(
                [chunk async for chunk in response.body_iterator]
            )

        workbook_bytes = asyncio.run(read_body())
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            [
                "SUMMARY REPORT",
                "Weekly Summary",
                "Daily Summary",
                "SETU",
                "KHT",
                "BLORA",
                "WUNUT",
            ],
        )

        summary = workbook["SUMMARY REPORT"]
        self.assertEqual(summary["B5"].value, 1000)
        self.assertEqual(summary["B6"].value, 350)
        self.assertEqual(summary["B7"].value, 340)
        self.assertEqual(summary["B8"].value, -660)

        weekly = workbook["Weekly Summary"]
        self.assertIn("target harian", weekly["A2"].value)
        self.assertEqual(weekly["C9"].value, 1000)
        self.assertEqual(weekly["D9"].value, 350)
        self.assertEqual(weekly["E9"].value, 340)

        daily = workbook["Daily Summary"]
        self.assertEqual(daily["D8"].value, 100)
        self.assertEqual(daily["E8"].value, 120)
        self.assertEqual(daily["D35"].value, 350)
        self.assertEqual(daily["E35"].value, 340)

        setu = workbook["SETU"]
        self.assertEqual(setu["D8"].value, "Customer FOB")
        self.assertEqual(setu["D9"].value, "TOTAL HARIAN")
        self.assertEqual(setu["E9"].value, 100)
        self.assertEqual(setu["F9"].value, 120)
        self.assertEqual(setu["D36"].value, "TOTAL BULAN")
        self.assertEqual(setu["E36"].value, 100)
        self.assertEqual(setu["F36"].value, 120)
        self.assertIn("attachment;", response.headers["content-disposition"])
        self.assertIn(".xlsx", response.headers["content-disposition"])

    def test_export_rejects_invalid_month(self):
        with self.assertRaises(main.HTTPException) as invalid:
            main.export_monthly_report(
                SimpleNamespace(session={"username": main.APP_USERNAME}),
                year=2026,
                month=13,
            )

        self.assertEqual(invalid.exception.status_code, 400)


if __name__ == "__main__":
    unittest.main()
